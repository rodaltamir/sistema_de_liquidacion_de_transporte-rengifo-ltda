import io
import re
from datetime import date, datetime
from typing import Dict, Any, List, Optional, Tuple
from dateutil import parser as date_parser
import openpyxl
from sqlalchemy.orm import Session
from sqlalchemy import func
from app.models.tenant import Viaje, ParametroLiquidacion, UnidadTransporte, Cliente, EmpresaApoyo
from app.services.calculation_engine import calculate_viaje_values, determine_trip_period


# -------------------------------------------------------------------------
# PALABRAS CLAVE PARA CLASIFICACIÓN INTELIGENTE DE CLIENTES Y TRANSPORTISTAS
# -------------------------------------------------------------------------
KNOWN_CLIENT_KEYWORDS = [
    "ACEITE", "SOYA", "FINO", "YPFB", "Y.P.F.B", "CARGILL", "BUNGE", "GRAVETAL",
    "ACEROS AREQUIPA", "ACEROSPAURITO", "AREQUIPA", "MINERA", "MINERIA", "FERTILIZANTES", "FERTIBOL", "FERTIPRO",
    "ALICONSUMO", "PIL ANDINA", "REPSOL", "PETROBRAS", "SHELL", "TOTAL", "COBOCE", "SOBOCE",
    "ITAGROBOL", "GREEN METALS", "GREENFIELD", "MINERA DON VALERIANO", "INDUSTRIA Y COMERCIO",
    "AGRO", "CEMENTO", "ALIMENTOS", "COMERCIALIZADORA", "CORPORACION MINERA", "CAMPOSOL", "PRODECO",
    "HAMMER", "LAMBOL", "PLUS METALS", "SYNTEX", "DON VALERIANO", "METALS", "ASPHALT",
    "TEMPLARIOS", "TRUCKS TEMPLARIOS", "REALTRUCK", "REALTRUCKS", "BOCKTRANS", "MEDLOG",
    "LIMACHI", "PEDRO WEILLIAM"
]

KNOWN_TRANSPORT_KEYWORDS = [
    "TRANS ", "TRANS.", "TRANSPORTE", "TRANSPORTES",
    "TRAILER", "LOGISTICA", "LOGISTIC", "COOP. TRANS", "COOPERATIVA", "LINEA", "CARGO", "FLOTA",
    "DAYTEL", "MGTRAILER", "CITRANSCO", "ROSESTAR", "BRITANIC",
    "REMOLINO", "QUELION", "PETRO FRAYER", "HANDEK", "COLMENA",
    "FUENTES", "BRUMEBOL", "JACHI", "JARM", "TEMPESTAD", "TRUKSPER", "UNITED"
]

KNOWN_INDIVIDUAL_CARRIERS = [
    "ROSMERY QUISPE", "GONZAYO"
]


def classify_entity_name(name: str) -> str:
    """
    Clasifica inteligentemente si una entidad es un CLIENTE que contrató el servicio
    o una EMPRESA DE TRANSPORTE / Transportista asociado.
    """
    if not name:
        return "cliente"
    n_up = str(name).strip().upper()

    # 1. Transportistas conocidos y palabras clave de transporte
    for kw in KNOWN_TRANSPORT_KEYWORDS:
        if kw in n_up:
            return "empresa"

    for ic in KNOWN_INDIVIDUAL_CARRIERS:
        if ic in n_up:
            return "empresa"

    # 2. Clientes conocidos
    for kw in KNOWN_CLIENT_KEYWORDS:
        if kw in n_up:
            return "cliente"

    # 3. Entidades comerciales/industriales comunes (S.A., S.R.L., SRL, LTDA, SAS, INC.)
    if any(s in n_up for s in ("S.A.", "S.A", "S.R.L.", "SRL", "LTDA", "LIMITADA", "SAS", "INC.")):
        return "cliente"

    return "cliente"


def sanitize_schema_name(name: str) -> str:
    """Sanitiza el nombre de la empresa para usarlo como esquema PostgreSQL."""
    slug = re.sub(r'[^a-zA-Z0-9]+', '_', name).strip('_').lower()
    return f"empresa_{slug[:30]}"


def normalize_company_name(name: Any) -> str:
    """Normaliza nombres de empresas (apóstrofes, espacios múltiples, mayúsculas)."""
    if not name:
        return ""
    s = str(name).strip().replace("´", "").replace("`", "").replace("’", "").replace("'", "")
    s = re.sub(r'\s+', ' ', s)
    return s.upper()


def get_excel_sheets(file_bytes: bytes) -> List[str]:
    """Retorna los nombres de las hojas en el libro Excel de manera ultra rápida."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheets = list(wb.sheetnames)
    wb.close()
    return sheets


def parse_date_value(val: Any) -> Optional[date]:
    """Parsea valores de fecha de forma robusta soportando datetime, date, str y números de serie Excel."""
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, (int, float)):
        # Número de serie de Excel (días desde 1899-12-30)
        try:
            from datetime import timedelta
            return date(1899, 12, 30) + timedelta(days=int(val))
        except Exception:
            return None
    val_str = str(val).strip()
    if not val_str or val_str.upper() in ("NONE", "NULL", "FALTA", "-", "0"):
        return None
    try:
        dt = date_parser.parse(val_str, dayfirst=True)
        return dt.date()
    except Exception:
        try:
            dt = date_parser.parse(val_str, dayfirst=False)
            return dt.date()
        except Exception:
            return None


def clean_float(val: Any, default: float = 0.0) -> float:
    """Convierte un valor a float limpiando caracteres como comas y espacios."""
    if val is None or val == "":
        return default
    if isinstance(val, (int, float)):
        return float(val)
    val_str = str(val).strip().replace(" ", "").replace(",", ".")
    try:
        return float(val_str)
    except Exception:
        return default


def clean_string(val: Any) -> str:
    """Limpia cadenas de texto."""
    if val is None:
        return ""
    return str(val).strip()


def detect_columns_and_header_row(rows_iter) -> Tuple[Optional[int], Dict[str, int]]:
    """
    Escanea las primeras 15 filas para encontrar el encabezado y mapear dinámicamente las columnas.
    Soporta formato Driver S.R.L. / Plantilla de Importación y formato oficial del sistema.
    """
    col_map: Dict[str, int] = {}
    header_row_idx: Optional[int] = None

    for r_idx, row in enumerate(rows_iter, 1):
        if r_idx > 15:
            break
        row_vals = [clean_string(c) for c in row]
        row_upper = [v.upper() for v in row_vals]

        has_placa = any("PLACA" in v for v in row_upper)
        has_prod = any("PRODUCTO" in v for v in row_upper)
        has_vol = any("VOLUMEN" in v or "VOL" in v or "DESPACHADO" in v for v in row_upper)
        has_emp = any("EMPRESA" in v or "TRANSPORTE" in v or "CLIENTE" in v for v in row_upper)

        if has_placa and (has_prod or has_vol or has_emp):
            header_row_idx = r_idx
            for c_idx, val in enumerate(row_upper):
                if not val:
                    continue
                if "PLACA" in val and "placa" not in col_map:
                    col_map["placa"] = c_idx
                elif ("EMPRESA TRANSPORTE" in val or "TRANSPORTISTA" in val or "ASOCIADO" in val or "SOCIO" in val or "PROPIETARIO" in val or ("EMPRESA" in val and "CLIENTE" not in val and "FACTURA" not in val)) and "empresa" not in col_map:
                    col_map["empresa"] = c_idx
                elif ("CLIENTE" in val or "CONSIGNATARIO" in val or "REMITENTE" in val or "FACTURA EMITIDA" in val or "RAZON SOCIAL" in val or "RAZÓN SOCIAL" in val or "DESTINATARIO" in val or "PAGADOR" in val) and "cliente" not in col_map:
                    col_map["cliente"] = c_idx
                elif "PRODUCTO" in val and "producto" not in col_map:
                    col_map["producto"] = c_idx
                elif ("PLANTA DE CARGA" in val or "PLANTA CARGA" in val or ("ORIGEN" in val and "FECHA" not in val and "VOL" not in val)) and "planta_carga" not in col_map:
                    col_map["planta_carga"] = c_idx
                elif ("PLANTA DE DESCARGA" in val or "PLANTA DESCARGA" in val or ("DESTINO" in val and "FECHA" not in val and "VOL" not in val)) and "planta_descarga" not in col_map:
                    col_map["planta_descarga"] = c_idx
                elif "TRAMO" in val and "tramo" not in col_map:
                    col_map["tramo"] = c_idx
                elif ("CARGA" in val and "FECHA" in val) or val in ("FECHA DE CARGA", "F. CARGA", "FECHA CARGA", "F_CARGA"):
                    col_map["fecha_carga"] = c_idx
                elif ("RECEPCI" in val and "FECHA" in val) or ("DESCARGA" in val and "FECHA" in val) or val in ("FECHA RECEPCIÓN", "FECHA RECEPCION", "FECHA DESCARGA", "F. DESCARGA"):
                    col_map["fecha_descarga"] = c_idx
                elif ("DESPACHADO" in val and ("VOL" in val or "15" in val or "LITRO" in val)) or ("ORIGEN" in val and "VOL" in val):
                    col_map["vol_origen"] = c_idx
                elif ("RECEPCIONAD" in val and ("VOL" in val or "15" in val or "LITRO" in val)) or ("DESCARGADO" in val and "VOL" in val):
                    col_map["vol_recepcion"] = c_idx
                elif ("FACTURAR" in val or "FACTURADO" in val) and ("VOL" in val or "M3" in val or "LITRO" in val):
                    col_map["vol_facturado"] = c_idx
                elif "TARIFA" in val or ("FLETE" in val and ("BS" in val or "M3" in val or "US" in val) and "IMPORTE" not in val):
                    col_map["tarifa"] = c_idx
                elif "IMPORTE" in val and ("FACTURAR" in val or "TOTAL" in val or "BS" in val):
                    col_map["importe_total"] = c_idx
                elif ("N°" in val or "NRO" in val or "NUM" in val or "Nº" in val) and "FACTURA" in val and "fecha" not in val:
                    col_map["nro_factura"] = c_idx
                elif "FECHA" in val and "FACTURA" in val:
                    col_map["fecha_factura"] = c_idx
                elif ("MIC" in val or "DTA" in val or "MCI" in val) and "FECHA" not in val:
                    col_map["mic_dta"] = c_idx
                elif ("MIC" in val or "DTA" in val) and "FECHA" in val:
                    col_map["fecha_mic"] = c_idx
                elif "CONTRATO" in val:
                    col_map["contrato"] = c_idx
                elif "MERMA" in val and "DESCONTAR" in val and "BS" in val:
                    col_map["merma_descontar"] = c_idx
                elif "PRECIO MERMA" in val:
                    col_map["precio_merma"] = c_idx
                elif "LOTE" in val:
                    col_map["lote"] = c_idx
                elif "OBS" in val or "OBSERVACI" in val:
                    col_map["obs"] = c_idx
            break

    return header_row_idx, col_map


def parse_excel_viajes_data(
    file_bytes: bytes,
    sheet_name: Optional[str] = None,
    session: Optional[Session] = None,
    db: Optional[Session] = None,
    asociacion_id: Optional[int] = None,
    max_preview_records: Optional[int] = None,
    target_empresa_name: Optional[str] = None
) -> Dict[str, Any]:
    """
    Lee, valida y procesa una hoja Excel de viajes.
    Soporta formato por Empresa (individual) y formato por Asociación (multi-empresa).
    Maneja con total precisión:
    - Clasificación inteligente de Clientes vs Empresas de Transporte.
    - Servicios y fletes con "SIN PLACA" sin omitirlos ni distorsionar importes.
    - Conciliación de facturas de múltiples despachos con sus filas de subtotal.
    - Cuadre exacto de importes a facturar y volúmenes en m³ y litros.
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheets = list(wb.sheetnames)

    # Determinación de hojas a procesar
    if not sheet_name or sheet_name in ("__ALL__", "ALL", "TODAS", "todas"):
        target_sheets = sheets
        selected_sheet_display = "__ALL__"
    elif sheet_name in sheets:
        target_sheets = [sheet_name]
        selected_sheet_display = sheet_name
    elif "," in sheet_name:
        target_sheets = [s.strip() for s in sheet_name.split(",") if s.strip() in sheets]
        if not target_sheets:
            target_sheets = sheets
        selected_sheet_display = sheet_name
    else:
        target_sheets = [sheets[0]]
        selected_sheet_display = sheets[0]

    valid_records: List[Dict[str, Any]] = []
    skipped_records: List[Dict[str, Any]] = []
    periodos_detectados = set()
    placas_detectadas = set()
    hojas_procesadas_stats: List[Dict[str, Any]] = []

    # Diccionario para conciliación de facturas: {nro_factura: [trip_records]} y {nro_factura: subtotal_row}
    facturas_trips: Dict[str, List[Dict[str, Any]]] = {}
    facturas_subtotales: Dict[str, Dict[str, Any]] = {}

    total_vol_facturado_lt = 0.0
    total_flete_bs = 0.0
    global_fila_counter = 0

    for sheet_title in target_sheets:
        ws = wb[sheet_title]

        # Extraer año si la hoja tiene el año en su nombre (ej. V-2024 DRIVER, 2025, 2026)
        year_match = re.search(r'\b(20\d\d)\b', sheet_title)
        sheet_year = int(year_match.group(1)) if year_match else None

        # Detectar posible nombre de empresa transportista principal en el título de la hoja (ej. "DRIVER" -> "DRIVER S.R.L.")
        sheet_upper = sheet_title.upper()
        detected_sheet_company = None
        if "DRIVER" in sheet_upper:
            detected_sheet_company = "DRIVER S.R.L."
        elif target_empresa_name:
            detected_sheet_company = target_empresa_name

        # Tomar primeras 15 filas para detección de encabezados
        sample_rows = []
        for r in ws.iter_rows(values_only=True):
            sample_rows.append(r)
            if len(sample_rows) >= 15:
                break

        if not sample_rows:
            continue

        # Si no se detectó en el título de la hoja, revisar los textos de los encabezados para detectar la empresa transportista
        if not detected_sheet_company:
            for s_row in sample_rows:
                for c_cell in s_row:
                    c_cell_str = str(c_cell or "").upper()
                    if "DRIVER" in c_cell_str:
                        detected_sheet_company = "DRIVER S.R.L."
                        break
                if detected_sheet_company:
                    break

        header_row_idx, col_map = detect_columns_and_header_row(sample_rows)

        if not header_row_idx or "placa" not in col_map:
            # Fallback a posiciones estándar
            if ws.max_column and ws.max_column >= 18:
                header_row_idx = 4
                col_map = {
                    "fecha_carga": 1,
                    "fecha_descarga": 2,
                    "mic_dta": 3,
                    "placa": 4,
                    "tramo": 5,
                    "cliente": 6,
                    "nro_factura": 7,
                    "fecha_factura": 8,
                    "producto": 9,
                    "vol_origen": 10,
                    "vol_recepcion": 11,
                    "vol_facturado": 16,
                    "tarifa": 17
                }
            elif ws.max_column and ws.max_column >= 7:
                header_row_idx = 1
                col_map = {
                    "empresa": 0,
                    "producto": 1,
                    "planta_carga": 2,
                    "planta_descarga": 4,
                    "placa": 5,
                    "fecha_carga": 6,
                    "fecha_descarga": 7,
                    "vol_origen": 8,
                    "vol_recepcion": 9,
                    "vol_facturado": 15,
                    "tarifa": 16,
                    "importe_total": 17,
                    "fecha_factura": 18,
                    "nro_factura": 19,
                    "mic_dta": 20
                }
            else:
                continue

        has_two_separate_entity_columns = (
            col_map.get("empresa") is not None and 
            col_map.get("cliente") is not None and 
            col_map["empresa"] != col_map["cliente"]
        )
        carrier_company = target_empresa_name or detected_sheet_company or "DRIVER S.R.L."

        sheet_valid_count = 0
        consecutive_empty = 0

        # Iterar filas de datos
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            if r_idx <= header_row_idx:
                continue

            # Verificación de fin de datos
            if not any(c is not None and str(c).strip() != "" for c in row):
                consecutive_empty += 1
                if consecutive_empty >= 20:
                    break
                continue
            consecutive_empty = 0

            def get_c(key: str) -> Any:
                idx = col_map.get(key)
                if idx is not None and idx < len(row):
                    return row[idx]
                return None

            placa_raw = clean_string(get_c("placa"))
            empresa_raw = clean_string(get_c("empresa"))
            cliente_raw = clean_string(get_c("cliente"))
            producto_raw = clean_string(get_c("producto"))
            planta_c = clean_string(get_c("planta_carga"))
            planta_d = clean_string(get_c("planta_descarga"))
            tramo_raw = clean_string(get_c("tramo"))

            vol_o_raw = clean_float(get_c("vol_origen"), 0.0)
            vol_r_raw = clean_float(get_c("vol_recepcion"), 0.0)
            vol_f_raw = clean_float(get_c("vol_facturado"), 0.0)
            tarifa_raw = clean_float(get_c("tarifa"), 0.0)
            importe_raw = clean_float(get_c("importe_total"), 0.0)

            nro_fac_raw = clean_string(get_c("nro_factura"))
            mic_raw = clean_string(get_c("mic_dta"))
            contrato_raw = clean_string(get_c("contrato"))

            f_carga_raw = get_c("fecha_carga")
            f_descarga_raw = get_c("fecha_descarga")
            f_fac_raw = get_c("fecha_factura")

            placa_upper = placa_raw.upper().strip()

            col_a_raw = cliente_raw or empresa_raw
            col_a_upper = col_a_raw.upper().strip()

            # -------------------------------------------------------------
            # 1. IDENTIFICACIÓN DE FILAS DE SUBTOTAL DE FACTURAS
            # -------------------------------------------------------------
            # En planillas como Driver S.R.L., las facturas con múltiples viajes
            # tienen filas hijas (despachos) y una fila de subtotal (sin placa ni empresa).
            is_subtotal_row = (
                (not placa_upper and not col_a_upper) or 
                placa_upper in ("TOTAL", "SUBTOTAL", "RESUMEN") or
                col_a_upper in ("TOTAL", "SUBTOTAL", "RESUMEN")
            )

            if is_subtotal_row:
                if nro_fac_raw and (importe_raw > 0 or vol_f_raw > 0):
                    fac_key = str(nro_fac_raw).strip()
                    facturas_subtotales[fac_key] = {
                        "fila": r_idx,
                        "hoja": sheet_title,
                        "subtotal_vol_m3": vol_f_raw,
                        "subtotal_importe_bs": importe_raw
                    }
                    skipped_records.append({
                        "fila": r_idx,
                        "hoja": sheet_title,
                        "motivo": f"Subtotal de Factura N° {fac_key} ({importe_raw:,.2f} Bs) conciliado con viajes hijos."
                    })
                else:
                    skipped_records.append({
                        "fila": r_idx,
                        "hoja": sheet_title,
                        "motivo": f"Fila de total o resumen ignorada ({placa_raw or col_a_raw or 'Fila vacía'})"
                    })
                continue

            # Si no tiene placa, ni empresa/cliente, ni importe, es fila en blanco
            if not placa_upper and not col_a_upper and importe_raw == 0 and vol_f_raw == 0:
                continue

            # -------------------------------------------------------------
            # 2. NORMALIZACIÓN DE PLACA Y SERVICIOS "SIN PLACA"
            # -------------------------------------------------------------
            # En transporte de soya, granos o servicios generales facturados,
            # la factura cubre un volumen global con "SIN PLACA". ¡NO DEBE OMITIRSE!
            is_sin_placa = not placa_upper or placa_upper in ("SIN PLACA", "SINPLACA", "FALTA PLACA", "FALTAPLACA", "S/P", "NONE", "NULL", "-")

            if is_sin_placa:
                clean_placa = "SIN PLACA"
            else:
                clean_placa = placa_upper.replace(" ", "")

            # -------------------------------------------------------------
            # 3. EMPRESA TRANSPORTISTA Y CLIENTE FACTURADO (COLUMNA A)
            # -------------------------------------------------------------
            # En libros de ventas de empresas de transporte (ej. DRIVER S.R.L.),
            # la Columna A (titulada 'FACTURA EMITIDA A:' o 'Empresa Transporte')
            # contiene los CLIENTES a quienes se les brindó el servicio y se les facturó.
            # Todas las cisternas/placas pertenecen a la empresa transportista que emitió las facturas.
            # En esta importación NO hay flotas de apoyo.
            if has_two_separate_entity_columns:
                clean_empresa = normalize_company_name(empresa_raw) or carrier_company
                clean_cliente = normalize_company_name(cliente_raw)
            else:
                clean_empresa = carrier_company
                clean_cliente = normalize_company_name(col_a_raw)

            # Inferencia inteligente de cliente ÚNICAMENTE si la celda de Columna A vino en blanco
            if not clean_cliente or clean_cliente in ("NONE", "NULL", "-", ""):
                prod_u = clean_string(producto_raw).upper()
                if any(x in prod_u for x in ("BOBINA", "ACERO", "FIERRO", "BARRA", "PLANCHA", "PERFIL")):
                    clean_cliente = "CORPORACION ACEROS AREQUIPA S.A."
                elif any(x in prod_u for x in ("SOYA", "GIRASOL", "HARINA", "COMPENSACION", "ACEITE")):
                    clean_cliente = "INDUSTRIAS DE ACEITE S.A."
                elif any(x in prod_u for x in ("FERTILIZANTE", "BORATO", "UREA")):
                    clean_cliente = "FERTILIZANTES DE LOS ANDES S.R.L."
                elif any(x in prod_u for x in ("MINERAL", "CONCENTRADO", "PLOMO", "ZINC", "ESTAÑO")):
                    clean_cliente = "PRODECO S.R.L."
                elif any(x in prod_u for x in ("MALTA", "CEBADA")):
                    clean_cliente = "BOCKTRANS S.R.L."
                elif any(x in prod_u for x in ("DIESEL", "GASOLINA", "GLP", "CRUDO", "PETROLEO", "INSUMOS", "IYA")):
                    clean_cliente = "YPFB"
                else:
                    clean_cliente = "CLIENTE GENERAL"

            tipo_entidad_asignada = "cliente"
            es_apoyo = False
            empresa_apoyo_nombre = None

            # -------------------------------------------------------------
            # 4. TRAMO Y PRODUCTO
            # -------------------------------------------------------------
            if planta_c or planta_d:
                orig = planta_c if planta_c and planta_c.upper() != "FALTA MIC" else ""
                dest = planta_d if planta_d and planta_d.upper() != "FALTA MIC" else ""
                if orig and dest:
                    clean_tramo = f"{orig} - {dest}".upper()
                elif orig:
                    clean_tramo = orig.upper()
                elif dest:
                    clean_tramo = dest.upper()
                else:
                    clean_tramo = tramo_raw.upper() if tramo_raw else "TRAMO NO ESPECIFICADO"
            else:
                clean_tramo = tramo_raw.upper() if tramo_raw else "TRAMO NO ESPECIFICADO"

            clean_prod = producto_raw.upper() if producto_raw else "DIESEL"

            # -------------------------------------------------------------
            # 5. FECHAS E INFERENCIA DE PERÍODO
            # -------------------------------------------------------------
            date_fac = parse_date_value(f_fac_raw)
            date_carga = parse_date_value(f_carga_raw)
            date_descarga = parse_date_value(f_descarga_raw)

            if not date_carga:
                if date_descarga:
                    date_carga = date_descarga
                elif date_fac:
                    date_carga = date_fac
                elif sheet_year:
                    date_carga = date(sheet_year, 1, 1)
                else:
                    date_carga = date.today()

            if not date_descarga:
                date_descarga = date_carga

            # Período
            if f_carga_raw and parse_date_value(f_carga_raw):
                periodo = determine_trip_period(date_carga, date_descarga)
            elif f_descarga_raw and parse_date_value(f_descarga_raw):
                periodo = determine_trip_period(date_descarga, None)
            elif date_fac:
                periodo = determine_trip_period(date_fac, None)
            elif sheet_year:
                periodo = f"{sheet_year}-01"
            else:
                periodo = determine_trip_period(date_carga, date_descarga)

            periodos_detectados.add(periodo)
            if clean_placa != "SIN PLACA":
                placas_detectadas.add(clean_placa)

            # -------------------------------------------------------------
            # 6. VOLÚMENES Y FLETES (ARITMÉTICA EXACTA)
            # -------------------------------------------------------------
            # En la planilla oficial, la columna es 'Volumen a facturar (m3)'
            # Por tanto, el valor está en METROS CÚBICOS (m³)
            vol_o = round(vol_o_raw, 1)
            vol_r = round(vol_r_raw, 1)

            if vol_f_raw > 0:
                vol_m3 = round(vol_f_raw, 4)
                vol_f_litros = round(vol_m3 * 1000.0, 1)
            elif vol_r > 0:
                vol_f_litros = vol_r
                vol_m3 = round(vol_r / 1000.0, 4)
            elif vol_o > 0:
                vol_f_litros = vol_o
                vol_m3 = round(vol_o / 1000.0, 4)
            else:
                vol_f_litros = 0.0
                vol_m3 = 0.0

            # Tarifa
            tarifa = tarifa_raw
            if tarifa <= 0:
                if importe_raw > 0 and vol_m3 > 0:
                    tarifa = round(importe_raw / vol_m3, 2)
                elif importe_raw > 0:
                    tarifa = importe_raw
                else:
                    tarifa = 392.00

            # Importe a Facturar (Flete Total Bs)
            # Si el Excel trae 'Importe a Facturar (Bs)', ese es el valor contractual sagrado
            if importe_raw > 0:
                flete_calculado = round(importe_raw, 2)
            elif vol_m3 > 0 and tarifa > 0:
                flete_calculado = round(vol_m3 * tarifa, 2)
            else:
                flete_calculado = 0.0

            tipo_tarifa = "BS_TOTAL_VIAJE" if (vol_m3 == 0 and flete_calculado > 0) else "BS_POR_M3"

            # -------------------------------------------------------------
            # 7. MIC/DTA Y CONTRATO
            # -------------------------------------------------------------
            clean_mic = mic_raw.upper().strip() if mic_raw and mic_raw.upper().strip() not in ("FALTA MIC", "SIN MIC", "NONE", "") else None
            clean_contrato = contrato_raw.strip() if contrato_raw and contrato_raw.strip() not in ("NONE", "NULL", "") else None

            # -------------------------------------------------------------
            # 8. VALIDACIÓN DE COMPLETITUD
            # -------------------------------------------------------------
            missing_fields = []
            if flete_calculado <= 0 and vol_m3 <= 0:
                missing_fields.append("Volumen e Importe en 0")

            is_incompleto = len(missing_fields) > 0
            incompleto_reason = ", ".join(missing_fields) if is_incompleto else None

            global_fila_counter += 1
            sheet_valid_count += 1

            record = {
                "fila": global_fila_counter,
                "hoja": sheet_title,
                "fila_hoja": r_idx,
                "empresa": clean_empresa,
                "cliente": clean_cliente,
                "entidad_origen": col_a_raw or clean_cliente,
                "tipo_entidad": tipo_entidad_asignada,
                "es_apoyo": es_apoyo,
                "empresa_apoyo_nombre": empresa_apoyo_nombre,
                "placa": clean_placa,
                "tramo": clean_tramo,
                "producto": clean_prod,
                "fecha_carga": date_carga.isoformat(),
                "fecha_descarga": date_descarga.isoformat(),
                "fecha_factura": date_fac.isoformat() if date_fac else None,
                "nro_factura": nro_fac_raw if nro_fac_raw else None,
                "mic_dta": clean_mic,
                "contrato": clean_contrato,
                "volumen_origen_litros": vol_o,
                "volumen_recepcionado_litros": vol_r,
                "volumen_facturado": vol_f_litros,
                "volumen_facturado_m3": round(vol_m3, 3),
                "tarifa_flete": tarifa,
                "tipo_tarifa": tipo_tarifa,
                "flete_total_bs": flete_calculado,
                "periodo_mes": periodo,
                "is_incompleto": is_incompleto,
                "incompleto_reason": incompleto_reason,
                "status": "INCOMPLETO" if is_incompleto else "LISTO",
                "included": True
            }

            valid_records.append(record)
            total_vol_facturado_lt += vol_f_litros
            total_flete_bs += flete_calculado

            if nro_fac_raw:
                fac_k = str(nro_fac_raw).strip()
                if fac_k not in facturas_trips:
                    facturas_trips[fac_k] = []
                facturas_trips[fac_k].append(record)

        hojas_procesadas_stats.append({
            "hoja": sheet_title,
            "viajes": sheet_valid_count,
            "anio": sheet_year
        })

    wb.close()

    preview_items = valid_records[:max_preview_records] if max_preview_records else valid_records
    valid_placas = [p for p in placas_detectadas if p and p not in ("SIN PLACA", "SINPLACA", "FALTAPLACA", "TOTAL")]

    # -------------------------------------------------------------
    # 9. CONCILIACIÓN DE FACTURAS CON SUS SUBTOTAIS DEL EXCEL
    # -------------------------------------------------------------
    facturas_conciliacion = []
    for fac_num, trip_list in facturas_trips.items():
        sum_vol_m3 = round(sum(t.get("volumen_facturado_m3", 0.0) for t in trip_list), 3)
        sum_imp_bs = round(sum(t.get("flete_total_bs", 0.0) for t in trip_list), 2)
        f_date = trip_list[0].get("fecha_factura") if trip_list else None
        f_cliente = trip_list[0].get("cliente") if trip_list else ""
        f_empresa = trip_list[0].get("empresa") if trip_list else ""

        sub = facturas_subtotales.get(fac_num)
        if sub:
            subtotal_excel_bs = round(float(sub.get("subtotal_importe_bs", 0.0)), 2)
            diferencia_bs = round(abs(sum_imp_bs - subtotal_excel_bs), 2)
            cuadra = diferencia_bs <= 0.05
            tiene_subtotal_fila = True
        else:
            subtotal_excel_bs = sum_imp_bs
            diferencia_bs = 0.0
            cuadra = True
            tiene_subtotal_fila = False

        facturas_conciliacion.append({
            "nro_factura": fac_num,
            "fecha_factura": f_date,
            "cliente": f_cliente,
            "empresa": f_empresa,
            "total_viajes": len(trip_list),
            "total_volumen_m3": sum_vol_m3,
            "total_importe_bs": sum_imp_bs,
            "subtotal_excel_bs": subtotal_excel_bs,
            "diferencia_bs": diferencia_bs,
            "cuadra": cuadra,
            "tiene_subtotal_fila": tiene_subtotal_fila
        })

    facturas_conciliacion.sort(key=lambda x: str(x["nro_factura"]))

    # -------------------------------------------------------------
    # 10. DETECCIÓN Y CLASIFICACIÓN DE EMPRESAS Y CLIENTES
    # -------------------------------------------------------------
    from app.models.public import Empresa
    existing_empresas = db.query(Empresa).all() if db else []
    emp_lookup = {normalize_company_name(e.name): e for e in existing_empresas}

    # Estadísticas de empresas transportistas detectadas
    empresas_stats: Dict[str, Dict[str, Any]] = {}
    clientes_stats: Dict[str, Dict[str, Any]] = {}

    for r in valid_records:
        cli_nom = r.get("cliente")
        if cli_nom:
            if cli_nom not in clientes_stats:
                clientes_stats[cli_nom] = {
                    "nombre": cli_nom,
                    "total_viajes": 0,
                    "total_volumen_m3": 0.0,
                    "total_flete_bs": 0.0
                }
            clientes_stats[cli_nom]["total_viajes"] += 1
            clientes_stats[cli_nom]["total_volumen_m3"] += r.get("volumen_facturado_m3", 0.0)
            clientes_stats[cli_nom]["total_flete_bs"] += r.get("flete_total_bs", 0.0)

        emp_nom = r.get("empresa")
        if emp_nom:
            if emp_nom not in empresas_stats:
                empresas_stats[emp_nom] = {
                    "nombre": emp_nom,
                    "total_viajes": 0,
                    "total_volumen_m3": 0.0,
                    "total_flete_bs": 0.0,
                    "placas": set()
                }
            empresas_stats[emp_nom]["total_viajes"] += 1
            empresas_stats[emp_nom]["total_volumen_m3"] += r.get("volumen_facturado_m3", 0.0)
            empresas_stats[emp_nom]["total_flete_bs"] += r.get("flete_total_bs", 0.0)
            if r.get("placa") and r["placa"] != "SIN PLACA":
                empresas_stats[emp_nom]["placas"].add(r["placa"])

    empresas_detectadas_list = []
    for emp_name, stats in sorted(empresas_stats.items()):
        existing = emp_lookup.get(emp_name)
        if existing:
            action = "reutilizar"
            empresa_id = existing.id
            schema_name = existing.schema_name
            already_in_asociacion = (existing.asociacion_id == asociacion_id) if asociacion_id else False
        else:
            action = "crear"
            empresa_id = None
            schema_name = sanitize_schema_name(emp_name)
            already_in_asociacion = False

        empresas_detectadas_list.append({
            "nombre": emp_name,
            "tipo_sugerido": "empresa",
            "action": action,
            "empresa_id": empresa_id,
            "schema_name": schema_name,
            "already_in_asociacion": already_in_asociacion,
            "total_viajes": stats["total_viajes"],
            "total_volumen_m3": round(stats["total_volumen_m3"], 3),
            "total_volumen_litros": round(stats["total_volumen_m3"] * 1000.0, 1),
            "total_flete_bs": round(stats["total_flete_bs"], 2),
            "cisternas_count": len(stats["placas"])
        })

    clientes_detectados_list = []
    for cli_name, stats in sorted(clientes_stats.items()):
        clientes_detectados_list.append({
            "nombre": cli_name,
            "tipo_sugerido": "cliente",
            "action": "registrar_cliente",
            "total_viajes": stats["total_viajes"],
            "total_volumen_m3": round(stats["total_volumen_m3"], 3),
            "total_volumen_litros": round(stats["total_volumen_m3"] * 1000.0, 1),
            "total_flete_bs": round(stats["total_flete_bs"], 2)
        })

    is_multi_empresa = len(empresas_stats) > 1

    return {
        "success": True,
        "sheets": sheets,
        "selected_sheet": selected_sheet_display,
        "hojas_procesadas": hojas_procesadas_stats,
        "total_hojas": len(sheets),
        "total_valid_rows": len(valid_records),
        "total_skipped_rows": len(skipped_records),
        "records": preview_items,
        "skipped_details": skipped_records[:100],
        "is_multi_empresa": is_multi_empresa,
        "empresas_detectadas": empresas_detectadas_list,
        "clientes_detectados": clientes_detectados_list,
        "facturas_conciliacion": facturas_conciliacion,
        "summary": {
            "total_viajes": len(valid_records),
            "viajes_listos": sum(1 for r in valid_records if not r.get("is_incompleto")),
            "viajes_incompletos": sum(1 for r in valid_records if r.get("is_incompleto")),
            "total_volumen_litros": round(total_vol_facturado_lt, 1),
            "total_volumen_m3": round(total_vol_facturado_lt / 1000.0, 3),
            "total_flete_bs": round(total_flete_bs, 2),
            "periodos_detectados": sorted(list(periodos_detectados)),
            "total_placas_unicas": len(valid_placas),
            "total_empresas_detectadas": len(empresas_detectadas_list),
            "total_clientes_detectados": len(clientes_detectados_list),
            "total_facturas_detectadas": len(facturas_conciliacion),
            "facturas_cuadran_100": all(f["cuadra"] for f in facturas_conciliacion),
            "total_hojas_procesadas": len(hojas_procesadas_stats)
        }
    }


def execute_import_viajes(
    records: List[Dict[str, Any]],
    session: Session,
    params: Optional[ParametroLiquidacion] = None,
    overwrite_duplicates: bool = False
) -> Dict[str, Any]:
    """
    Inserta o actualiza una lista de registros de viajes en la base de datos del esquema correspondiente.
    Garantiza:
    - Registro de viajes con 'SIN PLACA' usando una unidad de flota genérica activa.
    - Creación automática de Clientes en tenant.clientes si no existen previamente.
    - Creación automática de Empresas de Apoyo en tenant.empresas_apoyo si no existen.
    - Preservación exacta de Fletes Totales (Bs) y conciliación al centavo.
    """
    if not params:
        params = session.query(ParametroLiquidacion).first() or ParametroLiquidacion()

    viajes_creados = 0
    viajes_actualizados = 0
    viajes_omitidos = 0
    clientes_registrados = 0
    errores = []
    periodos_detectados = set()

    # Cache de unidades
    unidades_cache = {u.placa: u for u in session.query(UnidadTransporte).all()}

    # Cache de clientes en el esquema tenant
    clientes_cache = {c.nombre.upper().strip(): c for c in session.query(Cliente).all()}

    # Cache de empresas de apoyo
    apoyo_cache = {a.nombre.upper().strip(): a for a in session.query(EmpresaApoyo).all()}

    # Si la importación actual no contiene flota de apoyo, sanear proactivamente
    # cualquier viaje previo que haya quedado marcado erróneamente como apoyo
    if all(not rec.get("es_apoyo") for rec in records):
        stale_apoyo = session.query(Viaje).filter(Viaje.es_apoyo == True).all()
        for sv in stale_apoyo:
            if sv.empresa_apoyo_nombre:
                cli = session.query(Cliente).filter(Cliente.nombre.ilike(sv.empresa_apoyo_nombre.strip())).first()
                if not cli:
                    cli = Cliente(nombre=sv.empresa_apoyo_nombre.strip().upper(), is_active=True)
                    session.add(cli)
                    session.commit()
                    session.refresh(cli)
                sv.cliente = cli.nombre
            sv.es_apoyo = False
            sv.empresa_apoyo_nombre = None
            sv.empresa_apoyo_id = None
        if stale_apoyo:
            session.commit()

        from app.models.tenant import UnidadApoyo
        if session.query(UnidadApoyo).count() == 0:
            session.query(EmpresaApoyo).delete()
            session.commit()
            apoyo_cache = {}

    for rec in records:
        if rec.get("included") is False:
            continue
        try:
            placa_input = clean_string(rec.get("placa")).upper()
            if not placa_input or placa_input in ("SINPLACA", "FALTAPLACA", "SIN PLACA", "S/P", "NONE", "NULL"):
                placa = "SIN PLACA"
            else:
                placa = placa_input.replace(" ", "")

            # Evitar filas de resumen que se hayan filtrado
            if placa in ("TOTAL", "SUBTOTAL", "RESUMEN"):
                continue

            date_carga = parse_date_value(rec.get("fecha_carga")) or date.today()
            date_descarga = parse_date_value(rec.get("fecha_descarga")) or date_carga
            date_fac = parse_date_value(rec.get("fecha_factura"))
            mic_str = clean_string(rec.get("mic_dta")).upper() or None
            if mic_str in ("FALTA MIC", "SIN MIC", "NONE", ""):
                mic_str = None

            vol_o = clean_float(rec.get("volumen_origen_litros"), 0.0)
            vol_r = clean_float(rec.get("volumen_recepcionado_litros"), 0.0)
            vol_f_litros = clean_float(rec.get("volumen_facturado"), 0.0)
            vol_f_m3 = clean_float(rec.get("volumen_facturado_m3"), 0.0)

            if vol_f_litros <= 0 and vol_f_m3 > 0:
                vol_f_litros = round(vol_f_m3 * 1000.0, 1)

            tarifa = clean_float(rec.get("tarifa_flete"), 392.00)
            tipo_t = clean_string(rec.get("tipo_tarifa")) or "BS_POR_M3"
            tramo = clean_string(rec.get("tramo")) or "TRAMO NO ESPECIFICADO"
            cliente = clean_string(rec.get("cliente"))
            producto = clean_string(rec.get("producto")) or "DIESEL"

            if not cliente:
                prod_u = producto.upper()
                if any(x in prod_u for x in ("BOBINA", "ACERO", "FIERRO", "BARRA", "PLANCHA", "PERFIL")):
                    cliente = "CORPORACION ACEROS AREQUIPA S.A."
                elif any(x in prod_u for x in ("SOYA", "GIRASOL", "HARINA", "COMPENSACION", "ACEITE")):
                    cliente = "INDUSTRIAS DE ACEITE S.A."
                elif any(x in prod_u for x in ("FERTILIZANTE", "BORATO", "UREA")):
                    cliente = "FERTILIZANTES DE LOS ANDES S.R.L."
                elif any(x in prod_u for x in ("MINERAL", "CONCENTRADO", "PLOMO", "ZINC", "ESTAÑO")):
                    cliente = "PRODECO S.R.L."
                elif any(x in prod_u for x in ("MALTA", "CEBADA")):
                    cliente = "BOCKTRANS S.R.L."
                elif any(x in prod_u for x in ("DIESEL", "GASOLINA", "GLP", "CRUDO", "PETROLEO", "INSUMOS", "IYA")):
                    cliente = "YPFB"
                else:
                    cliente = "CLIENTE GENERAL"
            nro_fac = clean_string(rec.get("nro_factura")) or None
            flete_total_rec = clean_float(rec.get("flete_total_bs"), 0.0)
            es_apoyo = bool(rec.get("es_apoyo", False))
            emp_apoyo_nom = clean_string(rec.get("empresa_apoyo_nombre")) if es_apoyo else None

            # -------------------------------------------------------------
            # REGISTRO AUTOMÁTICO DE CLIENTE SI NO EXISTE
            # -------------------------------------------------------------
            if cliente:
                cli_key = cliente.upper().strip()
                if cli_key not in clientes_cache:
                    nuevo_cli = Cliente(
                        nombre=cli_key,
                        is_active=True
                    )
                    session.add(nuevo_cli)
                    session.commit()
                    session.refresh(nuevo_cli)
                    clientes_cache[cli_key] = nuevo_cli
                    clientes_registrados += 1

            # -------------------------------------------------------------
            # REGISTRO AUTOMÁTICO DE EMPRESA DE APOYO SI APLICA
            # -------------------------------------------------------------
            if es_apoyo and emp_apoyo_nom:
                ap_key = emp_apoyo_nom.upper().strip()
                if ap_key not in apoyo_cache:
                    nueva_apoyo = EmpresaApoyo(
                        nombre=ap_key,
                        is_active=True
                    )
                    session.add(nueva_apoyo)
                    session.commit()
                    session.refresh(nueva_apoyo)
                    apoyo_cache[ap_key] = nueva_apoyo

            # -------------------------------------------------------------
            # CREAR O RECUPERAR UNIDAD DE TRANSPORTE (INCLUYE "SIN PLACA")
            # -------------------------------------------------------------
            u = unidades_cache.get(placa)
            if not u:
                u = session.query(UnidadTransporte).filter(UnidadTransporte.placa == placa).first()
            if not u:
                u = UnidadTransporte(
                    placa=placa,
                    marca="Carga General / Sin Asignar" if placa == "SIN PLACA" else "Tractocamión / Cisterna",
                    capacidad_litros=vol_o if vol_o > 0 else 34000.0,
                    capacidad_m3=(vol_o / 1000.0) if vol_o > 0 else 34.0,
                    estado="Activo"
                )
                session.add(u)
                session.commit()
                session.refresh(u)
                unidades_cache[placa] = u

            calc = calculate_viaje_values(
                volumen_origen_litros=vol_o,
                volumen_recepcionado_litros=vol_r,
                producto=producto,
                tarifa_flete=tarifa,
                tipo_tarifa=tipo_t,
                params=params,
                volumen_facturado=vol_f_litros,
                flete_total_override=flete_total_rec if flete_total_rec > 0 else None
            )

            periodo = determine_trip_period(date_carga, date_descarga)
            periodos_detectados.add(periodo)

            # Buscar viaje existente para actualizar en lugar de duplicar
            existing_viaje = None
            if rec.get("id"):
                existing_viaje = session.query(Viaje).filter(Viaje.id == rec["id"]).first()
            elif nro_fac and placa and date_carga:
                existing_viaje = session.query(Viaje).filter(
                    Viaje.nro_factura == nro_fac,
                    Viaje.placa == placa,
                    Viaje.fecha_carga == date_carga
                ).first()
            elif overwrite_duplicates and placa and date_carga:
                existing_viaje = session.query(Viaje).filter(
                    Viaje.placa == placa,
                    Viaje.fecha_carga == date_carga
                ).first()

            if existing_viaje:
                existing_viaje.placa = placa
                existing_viaje.unidad_id = u.id
                existing_viaje.tramo = tramo
                existing_viaje.cliente = cliente
                existing_viaje.producto = producto
                existing_viaje.fecha_carga = date_carga
                existing_viaje.fecha_descarga = date_descarga
                existing_viaje.periodo_mes = periodo
                if nro_fac:
                    existing_viaje.nro_factura = nro_fac
                if date_fac:
                    existing_viaje.fecha_factura = date_fac
                if mic_str:
                    existing_viaje.mic_dta = mic_str
                existing_viaje.volumen_origen_litros = vol_o
                existing_viaje.volumen_recepcionado_litros = vol_r
                existing_viaje.volumen_facturado = calc["volumen_facturado"]
                existing_viaje.merma_real_litros = calc["merma_real_litros"]
                existing_viaje.tolerancia_pct = calc["tolerancia_pct"]
                existing_viaje.merma_tolerable_litros = calc["merma_tolerable_litros"]
                existing_viaje.merma_excedente_litros = calc["merma_excedente_litros"]
                existing_viaje.precio_merma_litro_bs = calc["precio_merma_litro_bs"]
                existing_viaje.merma_descontar_bs = calc["merma_descontar_bs"]
                existing_viaje.tarifa_flete = tarifa
                existing_viaje.tipo_tarifa = tipo_t
                existing_viaje.flete_total_bs = calc["flete_total_bs"]
                existing_viaje.es_apoyo = es_apoyo
                existing_viaje.empresa_apoyo_nombre = emp_apoyo_nom if es_apoyo else None
                existing_viaje.empresa_apoyo_id = None if not es_apoyo else existing_viaje.empresa_apoyo_id
                viajes_actualizados += 1
                continue

            nuevo_viaje = Viaje(
                mic_dta=mic_str,
                lote_codigo="1",
                unidad_id=u.id,
                placa=placa,
                tramo=tramo,
                cliente=cliente,
                nro_factura=nro_fac,
                fecha_factura=date_fac,
                producto=producto,
                fecha_carga=date_carga,
                fecha_descarga=date_descarga,
                periodo_mes=periodo,
                volumen_origen_litros=vol_o,
                volumen_recepcionado_litros=vol_r,
                volumen_facturado=calc["volumen_facturado"],
                merma_real_litros=calc["merma_real_litros"],
                tolerancia_pct=calc["tolerancia_pct"],
                merma_tolerable_litros=calc["merma_tolerable_litros"],
                merma_excedente_litros=calc["merma_excedente_litros"],
                precio_merma_litro_bs=calc["precio_merma_litro_bs"],
                merma_descontar_bs=calc["merma_descontar_bs"],
                tarifa_flete=tarifa,
                tipo_tarifa=tipo_t,
                flete_total_bs=calc["flete_total_bs"],
                es_apoyo=es_apoyo,
                empresa_apoyo_nombre=emp_apoyo_nom if es_apoyo else None,
                empresa_apoyo_id=None,
                estado="Pendiente"
            )
            session.add(nuevo_viaje)
            viajes_creados += 1
        except Exception as row_ex:
            errores.append(f"Fila {rec.get('fila', '?')}: {str(row_ex)}")

    if viajes_creados or viajes_actualizados:
        session.commit()

    return {
        "success": True,
        "total_importados": viajes_creados,
        "total_actualizados": viajes_actualizados,
        "total_omitidos": viajes_omitidos,
        "clientes_registrados": clientes_registrados,
        "periodos_detectados": sorted(list(periodos_detectados)),
        "errores": errores
    }


def execute_import_asociacion_multi_empresa(
    records: List[Dict[str, Any]],
    db: Session,
    asociacion_id: Optional[int] = None,
    nueva_asociacion_nombre: Optional[str] = None,
    nueva_asociacion_sigla: Optional[str] = None,
    overwrite_duplicates: bool = False,
    entidades_clasificacion: Optional[Dict[str, str]] = None,
    empresa_principal_nombre: Optional[str] = None
) -> Dict[str, Any]:
    """
    Importa despachos distribuyéndolos entre múltiples empresas y clientes.
    - Entidades clasificadas como 'cliente': sus viajes se asignan a la empresa principal
      de la asociación (ej. DRIVER S.R.L.) y el cliente se crea en tenant.clientes.
    - Entidades clasificadas como 'empresa': se crea o reutiliza su esquema dedicado en public.empresas.
    """
    from app.models.public import Asociacion, Empresa
    from app.db.session import get_tenant_session, init_tenant_schema
    from app.models.tenant import ParametroLiquidacion

    asoc = None
    if asociacion_id:
        asoc = db.query(Asociacion).filter(Asociacion.id == asociacion_id).first()
        if not asoc and not nueva_asociacion_nombre:
            raise ValueError(f"Asociación con ID {asociacion_id} no encontrada.")

    if not asoc and nueva_asociacion_nombre and nueva_asociacion_nombre.strip():
        asoc_name_clean = nueva_asociacion_nombre.strip()
        asoc = db.query(Asociacion).filter(Asociacion.name.ilike(asoc_name_clean)).first()
        if not asoc:
            asoc = Asociacion(
                name=asoc_name_clean,
                sigla=nueva_asociacion_sigla.strip().upper() if nueva_asociacion_sigla else None,
                is_active=True
            )
            db.add(asoc)
            db.commit()
            db.refresh(asoc)

    all_empresas = db.query(Empresa).all()
    emp_lookup = {normalize_company_name(e.name): e for e in all_empresas}

    # Determinar empresa anfitriona / principal para viajes de clientes
    main_empresa = None
    if empresa_principal_nombre:
        main_emp_norm = normalize_company_name(empresa_principal_nombre)
        main_empresa = emp_lookup.get(main_emp_norm)
    if not main_empresa and asoc:
        # Primera empresa de la asociación
        main_empresa = db.query(Empresa).filter(Empresa.asociacion_id == asoc.id).first()
    if not main_empresa and all_empresas:
        main_empresa = all_empresas[0]

    # Clasificaciones
    ent_map = entidades_clasificacion or {}

    # Agrupar registros válidos por empresa destino
    records_by_company: Dict[str, List[Dict[str, Any]]] = {}
    clientes_encontrados: set = set()

    for r in records:
        if r.get("included") is False:
            continue

        emp_raw = r.get("empresa")
        cli_raw = r.get("cliente")
        ent_origen = r.get("entidad_origen") or emp_raw or ""

        # Determinar si la entidad es un cliente o empresa
        ent_norm = normalize_company_name(ent_origen)
        clasificacion = ent_map.get(ent_norm) or r.get("tipo_entidad") or classify_entity_name(ent_norm)

        if clasificacion == "cliente":
            clientes_encontrados.add(ent_norm)
            r["cliente"] = ent_norm
            # Asignar viaje a la empresa principal
            if main_empresa:
                target_comp_name = normalize_company_name(main_empresa.name)
            else:
                target_comp_name = normalize_company_name(empresa_principal_nombre or "EMPRESA PRINCIPAL")
            r["empresa"] = target_comp_name
        else:
            target_comp_name = normalize_company_name(emp_raw) if emp_raw else "EMPRESA ASOCIADA"

        if target_comp_name not in records_by_company:
            records_by_company[target_comp_name] = []
        records_by_company[target_comp_name].append(r)

    if not records_by_company:
        return {
            "success": False,
            "message": "No hay registros seleccionados para importar.",
            "total_empresas_procesadas": 0,
            "empresas_creadas": 0,
            "empresas_reutilizadas": 0,
            "total_viajes_importados": 0,
            "total_viajes_actualizados": 0,
            "empresas_detalle": []
        }

    empresas_results = []
    total_importados_global = 0
    total_actualizados_global = 0
    total_creadas = 0
    total_reutilizadas = 0
    total_clientes_creados = 0

    for comp_name, comp_records in records_by_company.items():
        emp = emp_lookup.get(comp_name)
        is_new = False

        if emp:
            total_reutilizadas += 1
            if asoc and emp.asociacion_id != asoc.id:
                emp.asociacion_id = asoc.id
                db.commit()
                db.refresh(emp)
        else:
            is_new = True
            total_creadas += 1
            base_schema = sanitize_schema_name(comp_name)
            schema_name = base_schema
            counter = 1
            while db.query(Empresa).filter(Empresa.schema_name == schema_name).first():
                schema_name = f"{base_schema}_{counter}"
                counter += 1

            emp = Empresa(
                name=comp_name,
                schema_name=schema_name,
                nit=None,
                tipo_empresa="Sociedad",
                asociacion_id=asoc.id if asoc else None,
                icon="Truck",
                is_active=True
            )
            db.add(emp)
            db.commit()
            db.refresh(emp)
            emp_lookup[comp_name] = emp

            # Inicializar esquema PostgreSQL dedicado
            init_tenant_schema(emp.schema_name)

        # Importar viajes al esquema de la empresa
        tenant_session = get_tenant_session(emp.schema_name)
        try:
            params = tenant_session.query(ParametroLiquidacion).first()
            if not params:
                params = ParametroLiquidacion()
                tenant_session.add(params)
                tenant_session.commit()

            res = execute_import_viajes(
                records=comp_records,
                session=tenant_session,
                params=params,
                overwrite_duplicates=overwrite_duplicates
            )

            total_clientes_creados += res.get("clientes_registrados", 0)

            empresas_results.append({
                "empresa_id": emp.id,
                "empresa_name": emp.name,
                "schema_name": emp.schema_name,
                "accion": "creada" if is_new else "reutilizada",
                "total_viajes": len(comp_records),
                "importados": res["total_importados"],
                "actualizados": res["total_actualizados"],
                "omitidos": res["total_omitidos"],
                "clientes_registrados": res.get("clientes_registrados", 0),
                "errores": res["errores"]
            })
            total_importados_global += res["total_importados"]
            total_actualizados_global += res["total_actualizados"]
        finally:
            tenant_session.close()

    asoc_name_display = asoc.name if asoc else "Asociación"
    return {
        "success": True,
        "asociacion_id": asoc.id if asoc else None,
        "asociacion_name": asoc_name_display,
        "total_empresas_procesadas": len(records_by_company),
        "empresas_creadas": total_creadas,
        "empresas_reutilizadas": total_reutilizadas,
        "total_clientes_creados": total_clientes_creados,
        "total_viajes_importados": total_importados_global,
        "total_viajes_actualizados": total_actualizados_global,
        "empresas_detalle": empresas_results,
        "message": f"Se procesaron con éxito {len(records_by_company)} empresas ({total_creadas} creadas, {total_reutilizadas} reutilizadas) y {total_clientes_creados} clientes registrados con un total de {total_importados_global} despachos en {asoc_name_display}."
    }
