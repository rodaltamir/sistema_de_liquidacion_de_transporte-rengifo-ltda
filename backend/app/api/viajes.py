import io
import json
from datetime import date, datetime
from dateutil import parser as date_parser
import openpyxl
from fastapi import APIRouter, Depends, HTTPException, status, Query, UploadFile, File
from fastapi.responses import Response
from typing import List, Optional
from app.db.session import get_tenant_session
from app.api.deps import verify_tenant_exists
from app.models.public import Empresa
from app.models.tenant import Viaje, ParametroLiquidacion, UnidadTransporte
from app.schemas.viaje import ViajeCreate, ViajeUpdate, ViajeResponse, ViajeCalculoPreview
from app.services.calculation_engine import calculate_viaje_values, determine_trip_period
from app.services.excel_export import generate_viajes_template_excel

from sqlalchemy import func

router = APIRouter(prefix="/tenants/{schema_name}/viajes", tags=["Viajes y Despachos"])

def process_descuentos(raw_list):
    if not raw_list:
        return "[]", 0.0
    if isinstance(raw_list, str):
        try:
            parsed = json.loads(raw_list)
        except Exception:
            return "[]", 0.0
    else:
        parsed = raw_list

    clean_items = []
    total = 0.0
    if isinstance(parsed, list):
        for item in parsed:
            if hasattr(item, "model_dump"):
                d = item.model_dump()
            elif isinstance(item, dict):
                d = item
            else:
                continue
            concepto = str(d.get("concepto") or "").strip()
            try:
                monto = float(d.get("monto") or 0.0)
            except Exception:
                monto = 0.0
            if concepto and monto > 0:
                clean_items.append({"concepto": concepto, "monto": round(monto, 2)})
                total += monto
    return json.dumps(clean_items), round(total, 2)



@router.get("/periodos")
def get_periodos_disponibles(
    schema_name: str,
    empresa: Empresa = Depends(verify_tenant_exists)
):
    session = get_tenant_session(schema_name)
    try:
        results = session.query(
            Viaje.periodo_mes,
            func.count(Viaje.id).label("total_viajes"),
            func.sum(Viaje.volumen_recepcionado_litros).label("total_volumen"),
            func.sum(Viaje.flete_total_bs).label("total_flete")
        ).group_by(Viaje.periodo_mes).order_by(Viaje.periodo_mes.desc()).all()
        
        periodos = []
        for r in results:
            if r.periodo_mes:
                periodos.append({
                    "periodo_mes": r.periodo_mes,
                    "total_viajes": r.total_viajes,
                    "total_volumen_litros": float(r.total_volumen or 0),
                    "total_flete_bs": float(r.total_flete or 0)
                })
        return periodos
    finally:
        session.close()


@router.get("/", response_model=List[ViajeResponse])
def list_viajes(
    schema_name: str,
    periodo_mes: Optional[str] = None,
    anio: Optional[int] = None,
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
    placa: Optional[str] = None,
    estado: Optional[str] = None,
    producto: Optional[str] = None,
    tramo: Optional[str] = None,
    nro_factura: Optional[str] = None,
    facturado: Optional[str] = None,
    search: Optional[str] = None,
    orden: Optional[str] = None,
    empresa: Empresa = Depends(verify_tenant_exists)
):
    session = get_tenant_session(schema_name)
    try:
        query = session.query(Viaje)
        if periodo_mes:
            query = query.filter(Viaje.periodo_mes == periodo_mes)
        if anio:
            query = query.filter(Viaje.periodo_mes.startswith(f"{anio}-"))
        if fecha_desde:
            query = query.filter(Viaje.fecha_carga >= fecha_desde)
        if fecha_hasta:
            query = query.filter(Viaje.fecha_carga <= fecha_hasta)
        if placa:
            query = query.filter(Viaje.placa == placa.upper().strip())
        if estado:
            query = query.filter(Viaje.estado == estado)
        if producto and producto.upper().strip() != "TODOS":
            query = query.filter(Viaje.producto == producto.upper().strip())
        if tramo:
            query = query.filter(Viaje.tramo == tramo.upper().strip())
        if nro_factura:
            query = query.filter(Viaje.nro_factura.ilike(f"%{nro_factura.strip()}%"))
        if facturado == "si":
            query = query.filter(Viaje.nro_factura.isnot(None), Viaje.nro_factura != "")
        elif facturado == "no":
            query = query.filter((Viaje.nro_factura.is_(None)) | (Viaje.nro_factura == ""))
        if search:
            st = f"%{search.strip()}%"
            query = query.filter(
                (Viaje.nro_factura.ilike(st)) |
                (Viaje.mic_dta.ilike(st)) |
                (Viaje.placa.ilike(st)) |
                (Viaje.tramo.ilike(st)) |
                (Viaje.cliente.ilike(st)) |
                (Viaje.producto.ilike(st))
            )
            
        if orden in ("asc", "agrupado"):
            return query.order_by(Viaje.tramo.asc(), Viaje.producto.asc(), Viaje.fecha_carga.asc(), Viaje.id.asc()).all()
        return query.order_by(Viaje.fecha_carga.desc(), Viaje.id.desc()).all()
    finally:
        session.close()

@router.post("/preview", response_model=ViajeCalculoPreview)
def preview_calculo_viaje(
    schema_name: str,
    volumen_origen_litros: float,
    volumen_recepcionado_litros: float,
    producto: str,
    tarifa_flete: float,
    tipo_tarifa: Optional[str] = "BS_POR_M3",
    precio_merma_override: Optional[float] = None,
    volumen_facturado: Optional[float] = None,
    total_descuentos_adicionales: Optional[float] = 0.0,
    total_descuentos_mermas: Optional[float] = 0.0,
    empresa: Empresa = Depends(verify_tenant_exists)
):
    session = get_tenant_session(schema_name)
    try:
        params = session.query(ParametroLiquidacion).first()
        if not params:
            params = ParametroLiquidacion()

        calc = calculate_viaje_values(
            volumen_origen_litros=volumen_origen_litros,
            volumen_recepcionado_litros=volumen_recepcionado_litros,
            producto=producto,
            tarifa_flete=tarifa_flete,
            tipo_tarifa=tipo_tarifa,
            params=params,
            precio_merma_litro_override=precio_merma_override,
            volumen_facturado=volumen_facturado,
            total_descuentos_adicionales=total_descuentos_adicionales or 0.0,
            total_descuentos_mermas=total_descuentos_mermas or 0.0
        )
        return calc
    finally:
        session.close()

@router.post("/", response_model=ViajeResponse)
def create_viaje(
    schema_name: str,
    viaje_in: ViajeCreate,
    empresa: Empresa = Depends(verify_tenant_exists)
):
    session = get_tenant_session(schema_name)
    try:
        params = session.query(ParametroLiquidacion).first()
        if not params:
            params = ParametroLiquidacion()

        # Validar duplicados
        placa_clean = viaje_in.placa.upper().strip()
        mic_clean = viaje_in.mic_dta.strip().upper() if viaje_in.mic_dta and viaje_in.mic_dta.strip() else None

        if mic_clean:
            dup_mic = session.query(Viaje).filter(Viaje.mic_dta == mic_clean).first()
            if dup_mic:
                raise HTTPException(
                    status_code=400,
                    detail=f"Ya existe un viaje registrado con el MIC/DTA '{mic_clean}' (Placa: {dup_mic.placa}, Fecha de carga: {dup_mic.fecha_carga})"
                )

        # Validar duplicado exacto por placa y fechas de carga y descarga
        dup_viaje = session.query(Viaje).filter(
            Viaje.placa == placa_clean,
            Viaje.fecha_carga == viaje_in.fecha_carga,
            Viaje.fecha_descarga == viaje_in.fecha_descarga,
            Viaje.volumen_recepcionado_litros == viaje_in.volumen_recepcionado_litros
        ).first()
        if dup_viaje:
            raise HTTPException(
                status_code=400,
                detail=f"Ya existe un viaje idéntico registrado para la placa {placa_clean} en las fechas {viaje_in.fecha_carga} a {viaje_in.fecha_descarga} (MIC: {dup_viaje.mic_dta or 'S/N'})"
            )

        # Determinar periodo mensual (respetando valor explícito si fue provisto)
        if viaje_in.periodo_mes and viaje_in.periodo_mes.strip():
            periodo_determinado = viaje_in.periodo_mes.strip()
        else:
            periodo_determinado = determine_trip_period(viaje_in.fecha_carga, viaje_in.fecha_descarga)

        # Manejo de listas dinámicas de descuentos
        desc_adic_json, total_desc_adic = process_descuentos(viaje_in.descuentos_adicionales)
        desc_mermas_json, total_desc_mermas = process_descuentos(viaje_in.descuentos_mermas)

        # Volumen facturado (en Litros, por defecto igual a volumen recepcionado)
        vol_fact = viaje_in.volumen_facturado
        if vol_fact is None or vol_fact <= 0:
            vol_fact = round(viaje_in.volumen_recepcionado_litros, 1)
        elif vol_fact < 1000.0:
            vol_fact = round(vol_fact * 1000.0, 1)

        # Calcular mermas y flete automáticamente
        calc = calculate_viaje_values(
            volumen_origen_litros=viaje_in.volumen_origen_litros,
            volumen_recepcionado_litros=viaje_in.volumen_recepcionado_litros,
            producto=viaje_in.producto,
            tarifa_flete=viaje_in.tarifa_flete,
            tipo_tarifa=viaje_in.tipo_tarifa,
            params=params,
            precio_merma_litro_override=viaje_in.precio_merma_litro_bs,
            volumen_facturado=vol_fact,
            total_descuentos_adicionales=total_desc_adic,
            total_descuentos_mermas=total_desc_mermas
        )

        # Buscar unidad_id si existe por la placa
        unidad = session.query(UnidadTransporte).filter(UnidadTransporte.placa == placa_clean).first()
        unidad_id = unidad.id if unidad else viaje_in.unidad_id

        viaje = Viaje(
            mic_dta=mic_clean,
            lote_codigo=viaje_in.lote_codigo,
            unidad_id=unidad_id,
            placa=placa_clean,
            es_apoyo=viaje_in.es_apoyo or False,
            empresa_apoyo_id=viaje_in.empresa_apoyo_id,
            empresa_apoyo_nombre=viaje_in.empresa_apoyo_nombre,
            tramo=viaje_in.tramo,
            cliente=viaje_in.cliente or "YPFB",
            producto=viaje_in.producto.upper(),
            fecha_carga=viaje_in.fecha_carga,
            fecha_descarga=viaje_in.fecha_descarga,
            periodo_mes=periodo_determinado,
            nro_factura=viaje_in.nro_factura.strip() if viaje_in.nro_factura else None,
            fecha_factura=viaje_in.fecha_factura,
            volumen_origen_litros=viaje_in.volumen_origen_litros,
            volumen_recepcionado_litros=viaje_in.volumen_recepcionado_litros,
            volumen_facturado=calc["volumen_facturado"],
            merma_real_litros=calc["merma_real_litros"],
            tolerancia_pct=calc["tolerancia_pct"],
            merma_tolerable_litros=calc["merma_tolerable_litros"],
            merma_excedente_litros=calc["merma_excedente_litros"],
            precio_merma_litro_bs=calc["precio_merma_litro_bs"],
            merma_descontar_bs=calc["merma_descontar_bs"],
            descuentos_adicionales=desc_adic_json,
            descuentos_mermas=desc_mermas_json,
            total_descuentos_adicionales_bs=total_desc_adic,
            total_descuentos_mermas_bs=total_desc_mermas,
            tarifa_flete=viaje_in.tarifa_flete,
            tipo_tarifa=viaje_in.tipo_tarifa or "BS_POR_M3",
            flete_total_bs=calc["flete_total_bs"],
            estado="Pendiente",
            observaciones=viaje_in.observaciones
        )
        session.add(viaje)
        session.commit()
        session.refresh(viaje)
        return viaje
    finally:
        session.close()

@router.get("/plantilla-excel")
def download_plantilla_viajes(schema_name: str, empresa: Empresa = Depends(verify_tenant_exists)):
    excel_buffer = generate_viajes_template_excel(empresa_name=empresa.name)
    clean_emp = "".join([c if c.isalnum() else "_" for c in empresa.name])[:20]
    filename = f"Plantilla_Viajes_{clean_emp}.xlsx"
    return Response(
        content=excel_buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@router.post("/importar-excel")
async def importar_viajes_excel(
    schema_name: str,
    file: UploadFile = File(...),
    empresa: Empresa = Depends(verify_tenant_exists)
):
    if not file.filename.endswith(('.xlsx', '.xlsm', '.xltx')):
        raise HTTPException(status_code=400, detail="El archivo debe tener formato Excel (.xlsx)")

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(filename=io.BytesIO(contents), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo leer el archivo Excel: {str(e)}")

    ws = wb.active
    session = get_tenant_session(schema_name)
    try:
        params = session.query(ParametroLiquidacion).first()
        if not params:
            params = ParametroLiquidacion()

        viajes_creados = []
        errores = []
        periodos_detectados = set()

        # Detección inteligente de encabezados y mapeo de columnas
        col_map = {}
        header_row_idx = 4
        for r_cand in range(1, 8):
            row_vals = [str(ws.cell(row=r_cand, column=c).value or "").strip().upper() for c in range(1, ws.max_column + 1)]
            for col_i, val in enumerate(row_vals, 1):
                if ("CARGA" in val and "FECHA" in val) or val == "FECHA DE CARGA" or val == "F. CARGA":
                    col_map["fecha_carga"] = col_i
                elif ("DESCARGA" in val and "FECHA" in val) or val == "FECHA DE DESCARGA" or val == "F. DESCARGA":
                    col_map["fecha_descarga"] = col_i
                elif "MIC" in val or "DTA" in val or "MCI" in val:
                    col_map["mic_dta"] = col_i
                elif "PLACA" in val:
                    col_map["placa"] = col_i
                elif "TRAMO" in val:
                    col_map["tramo"] = col_i
                elif "CLIENTE" in val:
                    col_map["cliente"] = col_i
                elif ("NRO" in val or "N°" in val or "NUM" in val or "Nº" in val) and "FACTURA" in val:
                    col_map["nro_factura"] = col_i
                elif ("FECHA" in val or "F." in val) and ("FACTURA" in val or "FACTURACION" in val or "FACTURACIÓN" in val):
                    col_map["fecha_factura"] = col_i
                elif "FACTURA" in val and "nro_factura" not in col_map and "fecha_factura" not in col_map:
                    col_map["nro_factura"] = col_i
                elif "PRODUCTO" in val:
                    col_map["producto"] = col_i
                elif "ORIGEN" in val:
                    col_map["vol_origen"] = col_i
                elif "FACTURADO" in val and "VOL" in val:
                    col_map["vol_facturado"] = col_i
                elif "RECEPCIONADO" in val or "DESCARGADO" in val:
                    col_map["vol_rec"] = col_i
                elif "TARIFA" in val:
                    col_map["tarifa"] = col_i
                elif "OBS" in val:
                    col_map["obs"] = col_i
                elif "LOTE" in val:
                    col_map["lote"] = col_i

            if "placa" in col_map and ("vol_origen" in col_map or "vol_rec" in col_map):
                header_row_idx = r_cand
                break

        # Si no se detectaron encabezados por texto, aplicar mapeo por posición según cantidad de columnas
        if not col_map.get("placa"):
            if ws.max_column >= 18:
                # Formato oficial 19 columnas (con Factura, Fecha Factura, Volumen Facturado, sin col Empresa)
                col_map = {
                    "fecha_carga": 2,
                    "fecha_descarga": 3,
                    "mic_dta": 4,
                    "placa": 5,
                    "tramo": 6,
                    "cliente": 7,
                    "nro_factura": 8,
                    "fecha_factura": 9,
                    "producto": 10,
                    "vol_origen": 11,
                    "vol_rec": 12,
                    "vol_facturado": 17,
                    "tarifa": 18
                }
                header_row_idx = 4
            elif ws.max_column >= 15:
                # Formato previo 17 columnas de la liquidación física
                col_map = {
                    "fecha_carga": 2,
                    "fecha_descarga": 3,
                    "mic_dta": 4,
                    "placa": 6,
                    "tramo": 7,
                    "cliente": 8,
                    "producto": 9,
                    "vol_origen": 10,
                    "vol_rec": 11,
                    "tarifa": 16
                }
                header_row_idx = 4
            else:
                # Formato estándar 14 columnas
                col_map = {
                    "mic_dta": 1,
                    "placa": 2,
                    "tramo": 3,
                    "cliente": 4,
                    "producto": 5,
                    "fecha_carga": 6,
                    "fecha_descarga": 7,
                    "vol_origen": 8,
                    "vol_rec": 9,
                    "tarifa": 10,
                    "tipo_tarifa": 11,
                    "precio_merma": 12,
                    "lote": 13,
                    "obs": 14
                }
                header_row_idx = 4

        seen_mics_in_batch = set()
        seen_trips_in_batch = set()

        start_row = header_row_idx + 1
        for r in range(start_row, ws.max_row + 1):
            def get_val(key):
                idx = col_map.get(key)
                return ws.cell(row=r, column=idx).value if idx else None

            mic_dta = get_val("mic_dta")
            placa = get_val("placa")
            tramo = get_val("tramo")
            cliente = get_val("cliente")
            nro_fac = get_val("nro_factura")
            f_fac = get_val("fecha_factura")
            producto = get_val("producto")
            f_carga = get_val("fecha_carga")
            f_descarga = get_val("fecha_descarga")
            vol_origen = get_val("vol_origen")
            vol_rec = get_val("vol_rec")
            vol_fact = get_val("vol_facturado")
            tarifa = get_val("tarifa")
            tipo_tarifa = get_val("tipo_tarifa")
            precio_merma = get_val("precio_merma")
            lote = get_val("lote")
            obs = get_val("obs")

            if not placa and not vol_origen and not vol_rec:
                continue

            try:
                placa_str = str(placa or "").strip().upper()
                if not placa_str:
                    errores.append(f"Fila {r}: Placa no especificada")
                    continue

                tramo_str = str(tramo or "TRAMO NO ESPECIFICADO").strip().upper()
                cliente_str = str(cliente or "YPFB").strip().upper()
                producto_raw = str(producto or "DIESEL").strip().upper()
                if "GAS" in producto_raw:
                    producto_str = "GASOLINA"
                elif "IYA" in producto_raw or "ADIT" in producto_raw:
                    producto_str = "IYA"
                elif "SOYA" in producto_raw:
                    producto_str = "SOYA A GRANEL"
                elif "GIRASOL" in producto_raw:
                    producto_str = "GIRASOL A GRANEL"
                elif "ACERO" in producto_raw or "PLANCHA" in producto_raw:
                    producto_str = "PLANCHAS DE ACERO"
                elif "CRUDO" in producto_raw:
                    producto_str = "CRUDO"
                else:
                    producto_str = "DIESEL"

                def parse_date(v):
                    if isinstance(v, (datetime, date)):
                        return v.date() if isinstance(v, datetime) else v
                    if isinstance(v, str):
                        try:
                            return date_parser.parse(v).date()
                        except Exception:
                            return date.today()
                    return date.today()

                date_carga = parse_date(f_carga) if f_carga else date.today()
                date_descarga = parse_date(f_descarga) if f_descarga else date_carga
                date_fac = parse_date(f_fac) if f_fac else None

                vol_o = float(str(vol_origen or 0).replace(",", "."))
                vol_r = float(str(vol_rec or 0).replace(",", "."))
                vol_f = float(str(vol_fact).replace(",", ".")) if vol_fact else None
                tf = float(str(tarifa or 392.00).replace(",", ".")) if tarifa else 392.00
                tipo_t = str(tipo_tarifa or "BS_POR_M3").strip().upper()
                if "USD" in tipo_t:
                    tipo_t = "USD_POR_M3"
                else:
                    tipo_t = "BS_POR_M3"

                pm = float(str(precio_merma).replace(",", ".")) if precio_merma else None

                # Determinar periodo con holgura de 1-2 días
                periodo = determine_trip_period(date_carga, date_descarga)
                periodos_detectados.add(periodo)

                # Control anti-duplicados por MIC/DTA
                mic_str = str(mic_dta).strip().upper() if mic_dta else None
                if mic_str:
                    if mic_str in seen_mics_in_batch:
                        errores.append(f"Fila {r}: Viaje omitido por MIC/DTA duplicado en el archivo ({mic_str})")
                        continue
                    exist_mic = session.query(Viaje).filter(Viaje.mic_dta == mic_str).first()
                    if exist_mic:
                        errores.append(f"Fila {r}: Omitido por MIC/DTA '{mic_str}' ya registrado en el sistema (Placa {exist_mic.placa})")
                        continue
                    seen_mics_in_batch.add(mic_str)

                # Control anti-duplicados por placa, fechas y volumen
                trip_key = (placa_str, date_carga, date_descarga, vol_r)
                if trip_key in seen_trips_in_batch:
                    errores.append(f"Fila {r}: Viaje omitido por registro idéntico duplicado en el archivo ({placa_str}, {date_carga})")
                    continue
                exist_trip = session.query(Viaje).filter(
                    Viaje.placa == placa_str,
                    Viaje.fecha_carga == date_carga,
                    Viaje.fecha_descarga == date_descarga,
                    Viaje.volumen_recepcionado_litros == vol_r
                ).first()
                if exist_trip:
                    errores.append(f"Fila {r}: Omitido por existir ya en sistema ({placa_str}, {date_carga} a {date_descarga})")
                    continue
                seen_trips_in_batch.add(trip_key)

                # Asegurar que exista la unidad
                u = session.query(UnidadTransporte).filter(UnidadTransporte.placa == placa_str).first()
                if not u:
                    u = UnidadTransporte(
                        placa=placa_str,
                        marca="Tractocamión / Cisterna",
                        capacidad_litros=vol_o,
                        estado="Activo"
                    )
                    session.add(u)
                    session.commit()
                    session.refresh(u)

                # Calcular mermas y flete automáticamente
                calc = calculate_viaje_values(
                    volumen_origen_litros=vol_o,
                    volumen_recepcionado_litros=vol_r,
                    producto=producto_str,
                    tarifa_flete=tf,
                    tipo_tarifa=tipo_t,
                    params=params,
                    precio_merma_litro_override=pm,
                    volumen_facturado=vol_f
                )

                viaje = Viaje(
                    mic_dta=mic_str,
                    lote_codigo=str(lote).strip() if lote else "1",
                    unidad_id=u.id,
                    placa=placa_str,
                    tramo=tramo_str,
                    cliente=cliente_str,
                    nro_factura=str(nro_fac).strip() if nro_fac else None,
                    fecha_factura=date_fac,
                    producto=producto_str,
                    fecha_carga=date_carga,
                    fecha_descarga=date_descarga,
                    periodo_mes=periodo,
                    volumen_origen_litros=vol_o,
                    volumen_recepcionado_litros=vol_r,
                    volumen_facturado=calc["volumen_facturado"],
                    merma_real_litros=calc["merma_real_litros"],
                    tolerancia_pct=calc["tolerancia_pct"],
                    merma_tolerable_litros=calc["merma_tolerable_litros"],
                    merma_excedente_litros=calc["merma_excedente_litros"],
                    precio_merma_litro_bs=calc["precio_merma_litro_bs"],
                    merma_descontar_bs=calc["merma_descontar_bs"],
                    tarifa_flete=tf,
                    tipo_tarifa=tipo_t,
                    flete_total_bs=calc["flete_total_bs"],
                    estado="Pendiente",
                    observaciones=str(obs).strip() if obs else None
                )
                session.add(viaje)
                viajes_creados.append(viaje)

            except Exception as row_ex:
                errores.append(f"Fila {r}: {str(row_ex)}")

        if viajes_creados:
            session.commit()

        return {
            "success": True,
            "message": f"Se importaron {len(viajes_creados)} viajes con éxito.",
            "total_importados": len(viajes_creados),
            "periodos_detectados": sorted(list(periodos_detectados)),
            "errores": errores
        }
    finally:
        session.close()

@router.get("/{id}", response_model=ViajeResponse)
def get_viaje(schema_name: str, id: int, empresa: Empresa = Depends(verify_tenant_exists)):
    session = get_tenant_session(schema_name)
    try:
        viaje = session.query(Viaje).filter(Viaje.id == id).first()

        if not viaje:
            raise HTTPException(status_code=404, detail="Viaje no encontrado")
        return viaje
    finally:
        session.close()

@router.put("/{id}", response_model=ViajeResponse)
def update_viaje(
    schema_name: str,
    id: int,
    viaje_in: ViajeUpdate,
    empresa: Empresa = Depends(verify_tenant_exists)
):
    session = get_tenant_session(schema_name)
    try:
        viaje = session.query(Viaje).filter(Viaje.id == id).first()
        if not viaje:
            raise HTTPException(status_code=404, detail="Viaje no encontrado")

        params = session.query(ParametroLiquidacion).first() or ParametroLiquidacion()
        update_data = viaje_in.model_dump(exclude_unset=True)

        # Manejo de listas de descuentos en actualización
        if "descuentos_adicionales" in update_data:
            desc_adic_json, total_desc_adic = process_descuentos(update_data["descuentos_adicionales"])
            viaje.descuentos_adicionales = desc_adic_json
            viaje.total_descuentos_adicionales_bs = total_desc_adic
            del update_data["descuentos_adicionales"]
            if "total_descuentos_adicionales_bs" in update_data:
                del update_data["total_descuentos_adicionales_bs"]

        if "descuentos_mermas" in update_data:
            desc_mermas_json, total_desc_mermas = process_descuentos(update_data["descuentos_mermas"])
            viaje.descuentos_mermas = desc_mermas_json
            viaje.total_descuentos_mermas_bs = total_desc_mermas
            del update_data["descuentos_mermas"]
            if "total_descuentos_mermas_bs" in update_data:
                del update_data["total_descuentos_mermas_bs"]

        for field, val in update_data.items():
            setattr(viaje, field, val)

        # Volumen facturado (en Litros, por defecto igual a volumen recepcionado)
        vol_fact = viaje.volumen_facturado
        if vol_fact is None or vol_fact <= 0:
            vol_fact = round(viaje.volumen_recepcionado_litros, 1)
        elif vol_fact < 1000.0:
            vol_fact = round(vol_fact * 1000.0, 1)

        # Recalcular si cambiaron volúmenes, producto o tarifas
        calc = calculate_viaje_values(
            volumen_origen_litros=viaje.volumen_origen_litros,
            volumen_recepcionado_litros=viaje.volumen_recepcionado_litros,
            producto=viaje.producto,
            tarifa_flete=viaje.tarifa_flete,
            tipo_tarifa=viaje.tipo_tarifa,
            params=params,
            precio_merma_litro_override=viaje.precio_merma_litro_bs,
            volumen_facturado=vol_fact,
            total_descuentos_adicionales=viaje.total_descuentos_adicionales_bs,
            total_descuentos_mermas=viaje.total_descuentos_mermas_bs
        )
        viaje.volumen_facturado = calc["volumen_facturado"]
        viaje.merma_real_litros = calc["merma_real_litros"]
        viaje.tolerancia_pct = calc["tolerancia_pct"]
        viaje.merma_tolerable_litros = calc["merma_tolerable_litros"]
        viaje.merma_excedente_litros = calc["merma_excedente_litros"]
        viaje.precio_merma_litro_bs = calc["precio_merma_litro_bs"]
        viaje.merma_descontar_bs = calc["merma_descontar_bs"]
        viaje.flete_total_bs = calc["flete_total_bs"]
        if "periodo_mes" in update_data and update_data["periodo_mes"]:
            viaje.periodo_mes = str(update_data["periodo_mes"]).strip()
        elif not viaje.periodo_mes:
            viaje.periodo_mes = determine_trip_period(viaje.fecha_carga, viaje.fecha_descarga)

        session.commit()
        session.refresh(viaje)
        return viaje
    finally:
        session.close()

@router.delete("/{id}")
def delete_viaje(schema_name: str, id: int, empresa: Empresa = Depends(verify_tenant_exists)):
    session = get_tenant_session(schema_name)
    try:
        viaje = session.query(Viaje).filter(Viaje.id == id).first()
        if not viaje:
            raise HTTPException(status_code=404, detail="Viaje no encontrado")
        session.delete(viaje)
        session.commit()
        return {"message": "Viaje eliminado"}
    finally:
        session.close()


