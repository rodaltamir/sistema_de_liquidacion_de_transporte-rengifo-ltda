"""
Script de Auditoría y Verificación de Integridad de Datos
Compara el archivo Excel original contra los registros almacenados en PostgreSQL.
"""
import openpyxl
from collections import defaultdict
from app.db.session import get_tenant_session
from app.models.tenant import Viaje, Cliente, UnidadTransporte

EXCEL_PATH = "/app/exports/DRIVER S.R.L VENTAS.xlsx"

def run_comprehensive_audit():
    print("=" * 80)
    print("INICIANDO AUDITORÍA INTEGRAL: EXCEL VS POSTGRESQL")
    print(f"Archivo analizado: {EXCEL_PATH}")
    print("=" * 80)

    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    
    excel_trips = []
    excel_subtotals = defaultdict(list)
    excel_skipped = []
    
    total_excel_flete = 0.0
    total_excel_vol_m3 = 0.0
    
    excel_invoices_trips = defaultdict(list)
    excel_clients = set()
    excel_placas = set()

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sample_rows = []
        for r in ws.iter_rows(values_only=True):
            sample_rows.append(r)
            if len(sample_rows) >= 15:
                break
                
        header_idx = None
        col_map = {}
        for r_idx, row in enumerate(sample_rows, 1):
            row_vals = [str(c or "").strip().upper() for c in row]
            if any("PLACA" in v for v in row_vals):
                header_idx = r_idx
                for c_idx, val in enumerate(row_vals):
                    if "PLACA" in val and "placa" not in col_map:
                        col_map["placa"] = c_idx
                    elif ("EMPRESA" in val or "CLIENTE" in val or "FACTURA EMITIDA" in val) and "cliente" not in col_map:
                        col_map["cliente"] = c_idx
                    elif "PRODUCTO" in val and "producto" not in col_map:
                        col_map["producto"] = c_idx
                    elif ("ORIGEN" in val or "PLANTA DE CARGA" in val or "PLANTA CARGA" in val) and "origen" not in col_map:
                        col_map["origen"] = c_idx
                    elif ("DESTINO" in val or "PLANTA DE DESCARGA" in val or "PLANTA DESCARGA" in val) and "destino" not in col_map:
                        col_map["destino"] = c_idx
                    elif "TRAMO" in val and "tramo" not in col_map:
                        col_map["tramo"] = c_idx
                    elif ("CARGA" in val and "FECHA" in val) and "fecha_carga" not in col_map:
                        col_map["fecha_carga"] = c_idx
                    elif ("RECEPCI" in val or "DESCARGA" in val) and "FECHA" in val and "fecha_descarga" not in col_map:
                        col_map["fecha_descarga"] = c_idx
                    elif "FECHA" in val and "FACTURA" in val and "fecha_factura" not in col_map:
                        col_map["fecha_factura"] = c_idx
                    elif ("N°" in val or "NRO" in val or "NUM" in val) and "FACTURA" in val and "nro_factura" not in col_map:
                        col_map["nro_factura"] = c_idx
                    elif ("FACTURAR" in val or "FACTURADO" in val) and ("VOL" in val or "M3" in val or "LITRO" in val) and "vol_facturado" not in col_map:
                        col_map["vol_facturado"] = c_idx
                    elif "TARIFA" in val and "tarifa" not in col_map:
                        col_map["tarifa"] = c_idx
                    elif "IMPORTE" in val and ("FACTURAR" in val or "TOTAL" in val or "BS" in val) and "importe_total" not in col_map:
                        col_map["importe_total"] = c_idx
                break
                
        consecutive_empty = 0
        sheet_trips_count = 0
        sheet_subtotals_count = 0
        sheet_flete = 0.0
        sheet_vol_m3 = 0.0
        
        last_seen_client = None
        invoice_client_map = {}

        for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            if r_idx <= header_idx:
                continue
                
            if not any(c is not None and str(c).strip() != "" for c in row):
                consecutive_empty += 1
                if consecutive_empty >= 50:
                    break
                continue
            consecutive_empty = 0
            
            def get_val(key):
                idx = col_map.get(key)
                if idx is not None and idx < len(row):
                    return row[idx]
                return None
                
            placa = str(get_val("placa") or "").strip().upper()
            cliente_raw = str(get_val("cliente") or "").strip().upper()
            producto_raw = str(get_val("producto") or "").strip().upper()
            origen_raw = str(get_val("origen") or "").strip().upper()
            destino_raw = str(get_val("destino") or "").strip().upper()
            tramo_raw = str(get_val("tramo") or "").strip().upper()
            nro_fac = str(get_val("nro_factura") or "").strip()
            
            try:
                imp = float(str(get_val("importe_total") or 0).strip().replace(",", "."))
            except Exception:
                imp = 0.0
                
            try:
                vol_f = float(str(get_val("vol_facturado") or 0).strip().replace(",", "."))
            except Exception:
                vol_f = 0.0
                
            has_operational_fields = bool(
                placa or producto_raw or origen_raw or destino_raw or tramo_raw
            )
            
            is_subtotal = (
                (not has_operational_fields) or 
                placa in ("TOTAL", "SUBTOTAL", "RESUMEN") or
                cliente_raw in ("TOTAL", "SUBTOTAL", "RESUMEN")
            )
            
            if is_subtotal:
                if nro_fac and (imp > 0 or vol_f > 0):
                    excel_subtotals[nro_fac].append({
                        "hoja": sheet_name,
                        "fila": r_idx,
                        "importe_bs": imp,
                        "volumen_m3": vol_f
                    })
                    sheet_subtotals_count += 1
                else:
                    excel_skipped.append({"hoja": sheet_name, "fila": r_idx, "motivo": "Fila sin factura"})
                continue
                
            if cliente_raw:
                last_seen_client = cliente_raw
                if nro_fac:
                    invoice_client_map[nro_fac] = cliente_raw
                effective_client = cliente_raw
            elif nro_fac and nro_fac in invoice_client_map:
                effective_client = invoice_client_map[nro_fac]
            elif last_seen_client:
                effective_client = last_seen_client
            else:
                effective_client = "CLIENTE GENERAL"
                
            excel_clients.add(effective_client)
            
            clean_placa = "SIN PLACA" if (not placa or placa in ("SIN PLACA", "SINPLACA", "S/P", "NONE", "-")) else placa.replace(" ", "")
            if clean_placa != "SIN PLACA":
                excel_placas.add(clean_placa)
                
            tramo = f"{origen_raw} - {destino_raw}".strip(" -") if (origen_raw or destino_raw) else tramo_raw
            
            trip_info = {
                "hoja": sheet_name,
                "fila": r_idx,
                "nro_factura": nro_fac,
                "cliente": effective_client,
                "placa": clean_placa,
                "tramo": tramo,
                "importe_bs": imp,
                "volumen_m3": vol_f
            }
            
            excel_trips.append(trip_info)
            if nro_fac:
                excel_invoices_trips[nro_fac].append(trip_info)
                
            sheet_trips_count += 1
            sheet_flete += imp
            sheet_vol_m3 += vol_f
            
        total_excel_flete += sheet_flete
        total_excel_vol_m3 += sheet_vol_m3
        print(f"Hoja {sheet_name}: {sheet_trips_count} viajes reales, {sheet_subtotals_count} subtotales, Flete: Bs {sheet_flete:,.2f}")

    wb.close()
    
    session = get_tenant_session("empresa_driver_s_r_l")
    try:
        db_viajes = session.query(Viaje).all()
        db_clientes = session.query(Cliente).all()
        db_unidades = session.query(UnidadTransporte).all()
        
        db_viajes_count = len(db_viajes)
        db_flete_total = sum(v.flete_total_bs or 0.0 for v in db_viajes)
        
        db_invoices_trips = defaultdict(list)
        db_placas = set()
        db_clientes_set = set()
        
        for v in db_viajes:
            if v.nro_factura:
                db_invoices_trips[v.nro_factura.strip()].append(v)
            if v.placa and v.placa != "SIN PLACA":
                db_placas.add(v.placa)
            if v.cliente:
                db_clientes_set.add(v.cliente.upper().strip())
                
        print("\n" + "=" * 80)
        print("COMPARACIÓN FINAL:")
        print(f"  Viajes reales: Excel = {len(excel_trips)} | DB = {db_viajes_count} | Coincidencia: {'CORRECTO' if len(excel_trips) == db_viajes_count else 'ERROR'}")
        print(f"  Facturas: Excel = {len(excel_invoices_trips)} | DB = {len(db_invoices_trips)} | Coincidencia: {'CORRECTO' if len(excel_invoices_trips) == len(db_invoices_trips) else 'ERROR'}")
        print(f"  Flete Total: Excel = Bs {total_excel_flete:,.2f} | DB = Bs {db_flete_total:,.2f} | Diferencia = Bs {abs(db_flete_total - total_excel_flete):.2f}")
        print(f"  Flotas de apoyo en DB: {sum(1 for v in db_viajes if v.es_apoyo)} (debe ser 0)")
        print(f"  Viajes sin fecha en DB: {sum(1 for v in db_viajes if v.fecha_carga is None or v.fecha_descarga is None)} (debe ser 0)")
        print("=" * 80)
    finally:
        session.close()

if __name__ == "__main__":
    run_comprehensive_audit()
